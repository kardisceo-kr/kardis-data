"""세계 동향 지표(indicators.json) 자동 갱신.

GitHub Actions에서 매일 실행한다. AI를 쓰지 않고 공개 데이터만 내려받아 계산한다.
출처 하나가 실패해도 이전 값을 그대로 두고 errors에 기록한다(빈 화면 방지).

출처와 사용 조건 (2026-10-09 확인)
- 원/달러: ECB 유로 기준환율(KRW/EUR, USD/EUR)로 환산. ECB 통계는 출처 표기 시 상업적 재사용 허용.
- 미국 기준금리(실효연방기금금리, EFFR): 뉴욕 연방준비은행. 이용약관 고지 문구를 함께 표시해야 함.
- 브렌트유, 식량가격지수: 세계은행 Commodity Price Data(Pink Sheet), CC BY 4.0.
- 중점협력국 물가상승률: 세계은행 World Development Indicators(FP.CPI.TOTL.ZG), CC BY 4.0.
- 공여국 ODA 총액: OECD 발표치를 indicators_manual.json에 사람이(또는 월요일 예약 작업이) 적는다.
IMF 데이터는 상업적 재사용에 별도 허가가 필요하여 쓰지 않는다.
"""
import csv, io, json, re, sys, urllib.request
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "indicators.json"
MANUAL = ROOT / "indicators_manual.json"
KST = timezone(timedelta(hours=9))
TODAY = datetime.now(KST).date()
UA = {"User-Agent": "kardis-data indicators bot (+https://www.kardis.kr)"}


def get(url, binary=False, timeout=60):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        b = r.read()
    return b if binary else b.decode("utf-8", "replace")


def r2(x, n=2):
    return None if x is None else round(float(x), n)


# ---------- 1. 원/달러 (ECB) ----------
def krw_usd():
    start = (TODAY - timedelta(days=130)).isoformat()
    url = ("https://data-api.ecb.europa.eu/service/data/EXR/D.KRW+USD.EUR.SP00.A"
           f"?startPeriod={start}&format=csvdata")
    rows = list(csv.DictReader(io.StringIO(get(url))))
    by = {}
    for r in rows:
        if not r.get("OBS_VALUE"):
            continue
        by.setdefault(r["TIME_PERIOD"], {})[r["CURRENCY"]] = float(r["OBS_VALUE"])
    series = [[d, r2(v["KRW"] / v["USD"], 1)] for d, v in sorted(by.items()) if "KRW" in v and "USD" in v]
    if len(series) < 20:
        raise ValueError(f"ECB 관측치 부족 {len(series)}")
    series = series[-90:]
    return {
        "id": "krwusd", "label": "원/달러 환율", "unit": "원", "freq": "일간",
        "value": series[-1][1], "asof": series[-1][0], "series": series,
        "change": r2(series[-1][1] - series[0][1], 1), "change_from": series[0][0],
        "why": "달러 강세는 원화로 편성된 사업비의 현지 구매력과, 달러 부채가 많은 협력국의 상환 부담을 바꿉니다.",
        "source": "ECB 유로 기준환율로 환산(kardis 계산)", "source_url": "https://data.ecb.europa.eu/data/datasets/EXR",
        "note": "Source: ECB statistics.",
    }


# ---------- 2. 미국 실효연방기금금리 (뉴욕 연준) ----------
def effr():
    start = (TODAY - timedelta(days=400)).isoformat()
    url = ("https://markets.newyorkfed.org/api/rates/unsecured/effr/search.json"
           f"?startDate={start}&endDate={TODAY.isoformat()}")
    d = json.loads(get(url))
    pts = sorted((x["effectiveDate"], float(x["percentRate"])) for x in d.get("refRates", []) if x.get("percentRate") is not None)
    if len(pts) < 20:
        raise ValueError(f"EFFR 관측치 부족 {len(pts)}")
    # 그래프는 주 1점(매주 마지막 관측치)으로 줄인다
    weekly = {}
    for dt, v in pts:
        y, w, _ = date.fromisoformat(dt).isocalendar()
        weekly[(y, w)] = [dt, r2(v)]
    series = [weekly[k] for k in sorted(weekly)][-53:]
    series[-1] = [pts[-1][0], r2(pts[-1][1])]
    return {
        "id": "effr", "label": "미국 기준금리(실효)", "unit": "%", "freq": "일간",
        "value": r2(pts[-1][1]), "asof": pts[-1][0], "series": series,
        "change": r2(series[-1][1] - series[0][1]), "change_from": series[0][0],
        "why": "미국 금리는 협력국의 외채 이자와 자본 유출입, 개발은행 차관 금리 환경에 영향을 줍니다.",
        "source": "뉴욕 연방준비은행 EFFR", "source_url": "https://www.newyorkfed.org/markets/reference-rates/effr",
        "note": "EFFR data is subject to the Terms of Use posted at newyorkfed.org. The New York Fed is not responsible for publication of the EFFR by kardis.",
    }


# ---------- 3·4. 브렌트유, 식량가격지수 (세계은행 Pink Sheet) ----------
def pink_sheet():
    try:
        import openpyxl  # noqa
    except ImportError:
        raise RuntimeError("openpyxl 필요")
    import openpyxl
    page = get("https://www.worldbank.org/en/research/commodity-markets")
    m = re.search(r'https://thedocs\.worldbank\.org/[^"\'\s]+CMO-Historical-Data-Monthly\.xlsx', page)
    if not m:
        raise ValueError("Pink Sheet 월간 파일 링크를 찾지 못함")
    wb = openpyxl.load_workbook(io.BytesIO(get(m.group(0), binary=True, timeout=120)), read_only=True, data_only=True)

    def column(sheet_name, pattern):
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        hdr_i = col = None
        for i, row in enumerate(rows[:12]):
            for j, c in enumerate(row):
                if isinstance(c, str) and re.search(pattern, c, re.I):
                    hdr_i, col = i, j
                    break
            if col is not None:
                break
        if col is None:
            peek = [[c for c in r[:14] if c is not None] for r in rows[:10]]
            raise ValueError(f"{sheet_name}: '{pattern}' 열 없음. 머리 행: {peek}"[:900])
        out = []
        for row in rows[hdr_i + 1:]:
            k = row[0]
            if isinstance(k, str) and re.fullmatch(r"\d{4}M\d{2}", k.strip()):
                v = row[col]
                if isinstance(v, (int, float)):
                    out.append([f"{k[:4]}-{k[5:7]}", v])
        if len(out) < 12:
            raise ValueError(f"{sheet_name}: 관측치 부족")
        return out[-24:]

    sheets = wb.sheetnames
    price_sheet = next(s for s in sheets if "Monthly Prices" in s)
    index_sheet = next(s for s in sheets if "Monthly Indices" in s)
    brent = [[d, r2(v)] for d, v in column(price_sheet, r"^Crude oil, Brent")]
    food = [[d, r2(v, 1)] for d, v in column(index_sheet, r"^\s*(i?FOOD|Food)\b")]
    common = {"freq": "월간", "source": "세계은행 Commodity Price Data(Pink Sheet), CC BY 4.0",
              "source_url": "https://www.worldbank.org/en/research/commodity-markets", "note": ""}
    return [
        dict(common, id="brent", label="국제유가(브렌트)", unit="달러/배럴", value=brent[-1][1], asof=brent[-1][0],
             series=brent, change=r2(brent[-1][1] - brent[0][1]), change_from=brent[0][0],
             why="유가는 협력국의 수입 물가와 재정, 사업 현장의 운송·연료비를 움직입니다."),
        dict(common, id="food", label="식량가격지수", unit="2010=100", value=food[-1][1], asof=food[-1][0],
             series=food, change=r2(food[-1][1] - food[0][1], 1), change_from=food[0][0],
             why="식량 가격은 식량 수입국의 물가와 취약계층 생계, 농업·영양 사업의 목표치에 직접 영향을 줍니다."),
    ]


# ---------- 5. 중점협력국 물가상승률 (세계은행 WDI) ----------
def cpi25():
    cs = json.loads((ROOT / "countries.json").read_text(encoding="utf-8"))["countries"]
    iso = ";".join(c["iso3"] for c in cs)
    url = f"https://api.worldbank.org/v2/country/{iso}/indicator/FP.CPI.TOTL.ZG?format=json&mrnev=1&per_page=100"
    d = json.loads(get(url))
    name = {c["iso3"]: c["ko"] for c in cs}
    rows = [{"iso3": x["countryiso3code"], "ko": name.get(x["countryiso3code"], x["countryiso3code"]),
             "year": x["date"], "value": r2(x["value"], 1)}
            for x in (d[1] or []) if x.get("value") is not None]
    if len(rows) < 15:
        raise ValueError(f"WDI 물가 국가 수 부족 {len(rows)}")
    rows.sort(key=lambda r: -r["value"])
    return {
        "id": "cpi25", "label": "중점협력국 물가상승률", "unit": "%", "freq": "연간",
        "rows": rows, "asof": max(r["year"] for r in rows), "missing": sorted(set(name) - {r["iso3"] for r in rows}),
        "why": "물가가 높은 나라일수록 사업비 산정 단가와 목표치의 현실성을 자주 다시 봐야 합니다.",
        "source": "세계은행 World Development Indicators(FP.CPI.TOTL.ZG), CC BY 4.0",
        "source_url": "https://data.worldbank.org/indicator/FP.CPI.TOTL.ZG", "note": "나라별 가장 최근 연도 값",
    }


def main():
    prev = {}
    if OUT.exists():
        try:
            prev = {x["id"]: x for x in json.loads(OUT.read_text(encoding="utf-8")).get("items", [])}
        except Exception:
            prev = {}
    items, errors = {}, []
    for fn, ids in [(krw_usd, ["krwusd"]), (effr, ["effr"]), (pink_sheet, ["brent", "food"]), (cpi25, ["cpi25"])]:
        try:
            res = fn()
            for x in (res if isinstance(res, list) else [res]):
                x["fetched"] = TODAY.isoformat()
                items[x["id"]] = x
        except Exception as e:  # 실패하면 이전 값을 유지
            errors.append(f"{fn.__name__}: {type(e).__name__}: {e}"[:1000])
            for i in ids:
                if i in prev:
                    items[i] = dict(prev[i], stale=True)
    manual = json.loads(MANUAL.read_text(encoding="utf-8")) if MANUAL.exists() else {"items": []}
    order = ["krwusd", "effr", "brent", "food"]
    out = {
        "updated": TODAY.isoformat(),
        "note": "매일 GitHub Actions가 공개 데이터로 계산한다(AI 미사용). stale=true는 최근 수집에 실패해 이전 값을 보여 주는 항목.",
        "items": [items[i] for i in order if i in items] + manual.get("items", []) + ([items["cpi25"]] if "cpi25" in items else []),
        "errors": errors,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"items={len(out['items'])} errors={len(errors)}")
    for e in errors:
        print("ERR", e)


if __name__ == "__main__":
    main()
